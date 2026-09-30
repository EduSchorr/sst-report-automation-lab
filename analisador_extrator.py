from __future__ import annotations

import hashlib
import io
from typing import Any, Dict

import docx
import fitz
import openpyxl

def calculate_sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()

def extract_docx_details(data: bytes, filename: str = "") -> Dict[str, Any]:
    document = docx.Document(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    full_text = "\n".join(paragraphs)

    headings = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        style = paragraph.style.name if paragraph.style else ""
        if text and (
            "Heading" in style
            or "Título" in style
            or (len(text) < 100 and text.isupper())
        ):
            headings.append(text)

    tables = []
    for table in document.tables[:5]:
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip().replace("\n", " ") for cell in row.cells])
        tables.append(rows)

    image_count = sum(1 for rel in document.part.rels.values() if "image" in rel.target_ref.lower())

    return {
        "filename": filename,
        "text_sample": full_text[:4000],
        "full_text_len": len(full_text),
        "word_count": len(full_text.split()),
        "paragraphs_count": len(document.paragraphs),
        "headings": headings[:30],
        "tables_count": len(document.tables),
        "tables_sample": tables,
        "images_count": image_count,
        "num_pages": 1,
        "is_scanned": False,
    }

def extract_pdf_details(data: bytes, filename: str = "") -> Dict[str, Any]:
    document = fitz.open(stream=data, filetype="pdf")
    chunks = []
    images = 0
    has_tables = False

    for page in list(document)[:25]:
        chunks.append(page.get_text())
        images += len(page.get_images())
        try:
            tables = page.find_tables()
            has_tables = has_tables or bool(tables and tables.tables)
        except Exception:
            pass

    full_text = "\n".join(chunks)
    return {
        "filename": filename,
        "text_sample": full_text[:4000],
        "full_text_len": len(full_text),
        "word_count": len(full_text.split()),
        "paragraphs_count": len(full_text.split("\n\n")),
        "headings": [],
        "tables_count": int(has_tables),
        "tables_sample": [],
        "images_count": images,
        "num_pages": len(document),
        "is_scanned": len(full_text.strip()) < 100 and images > 0,
    }

def extract_xlsx_details(data: bytes, filename: str = "") -> Dict[str, Any]:
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    sheets = workbook.sheetnames
    text = "Planilha com abas: " + ", ".join(sheets)
    return {
        "filename": filename,
        "text_sample": text,
        "full_text_len": len(text),
        "word_count": len(text.split()),
        "paragraphs_count": len(sheets),
        "headings": sheets,
        "tables_count": len(sheets),
        "tables_sample": [],
        "images_count": 0,
        "num_pages": 1,
        "is_scanned": False,
    }
